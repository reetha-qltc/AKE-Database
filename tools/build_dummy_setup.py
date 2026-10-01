"""Build a DUMMY finance setup for an Indian SAP Business One test company - no Service Layer needed.

Run:  python tools/build_dummy_setup.py
Output (data/dummy/):
  Dummy_Finance_Setup.xlsx      - COA, GST tax codes, TDS codes, G/L determination (reference + manual entry sheet)
  dtw/COA_L2.csv ... COA_L4.csv - Chart of Accounts in DTW (Data Transfer Workbench) format, import in level order

Before importing, set DRAWERS below to the drawer codes of YOUR company (Level-1 accounts):
  SELECT "AcctCode", "AcctName" FROM OACT WHERE "Levels" = 1 ORDER BY "AcctCode";
then re-run this script.
"""
import csv, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "dummy"
DTW = OUT / "dtw"
DTW.mkdir(parents=True, exist_ok=True)

# Drawer (Level-1) account codes in the target company. EDIT THESE before importing via DTW.
# If the company has a single expense drawer, point COST_OF_SALES and EXPENSES to the same code.
DRAWERS = {
    "ASSETS": "<Assets drawer code>",
    "LIABILITIES": "<Liabilities drawer code>",
    "EQUITY": "<Capital and Reserves drawer code>",
    "REVENUE": "<Revenue / Turnover drawer code>",
    "COST_OF_SALES": "<Cost of Sales drawer code>",
    "EXPENSES": "<Operating Costs / Expenses drawer code>",
}
ACCT_TYPE = {"REVENUE": "at_Revenues", "COST_OF_SALES": "at_Expenses", "EXPENSES": "at_Expenses"}
CURRENCY = "INR"

# ---------------------------------------------------------------- Chart of Accounts (dummy)
# (code, name, children)  - an entry WITH children is a title, an entry without is postable.
COA = [
    ("ASSETS", [
        ("100000", "Fixed Assets", [
            ("100100", "Land and Building", None),
            ("100200", "Plant and Machinery", None),
            ("100300", "Furniture and Fixtures", None),
            ("100400", "Computers and Peripherals", None),
            ("100500", "Vehicles", None),
            ("100900", "Accumulated Depreciation", None),
        ]),
        ("110000", "Current Assets", [
            ("110100", "Inventory", [
                ("110101", "Raw Material Stock", None),
                ("110102", "Work in Progress", None),
                ("110103", "Finished Goods Stock", None),
                ("110104", "Stores and Spares", None),
            ]),
            ("110200", "Trade Receivables", [
                ("110201", "Domestic Debtors", None),
                ("110202", "Export Debtors", None),
            ]),
            ("110300", "Cash and Bank", [
                ("110301", "Cash in Hand", None),
                ("110302", "Petty Cash", None),
                ("110303", "HDFC Bank Current Account", None),
                ("110304", "ICICI Bank Current Account", None),
            ]),
            ("110400", "GST Input Credit", [
                ("110401", "Input CGST", None),
                ("110402", "Input SGST", None),
                ("110403", "Input IGST", None),
            ]),
            ("110500", "TDS and Advance Tax", [
                ("110501", "TDS Receivable (deducted by customers)", None),
                ("110502", "TCS Receivable", None),
                ("110503", "Advance Income Tax", None),
            ]),
            ("110600", "Loans and Advances", [
                ("110601", "Advance to Suppliers", None),
                ("110602", "Staff Advances", None),
                ("110603", "Security Deposits", None),
                ("110604", "Prepaid Expenses", None),
            ]),
        ]),
    ]),
    ("LIABILITIES", [
        ("200000", "Non-Current Liabilities", [
            ("200100", "Term Loan - HDFC Bank", None),
            ("200200", "Unsecured Loans from Directors", None),
        ]),
        ("210000", "Current Liabilities", [
            ("210100", "Trade Payables", [
                ("210101", "Domestic Creditors", None),
                ("210102", "Import Creditors", None),
            ]),
            ("210200", "GST Output Liability", [
                ("210201", "Output CGST", None),
                ("210202", "Output SGST", None),
                ("210203", "Output IGST", None),
            ]),
            ("210300", "TDS Payable", [
                ("210301", "TDS Payable - 194C Contractors", None),
                ("210302", "TDS Payable - 194J Professional and Technical", None),
                ("210303", "TDS Payable - 194I Rent", None),
                ("210304", "TDS Payable - 194H Commission", None),
                ("210305", "TDS Payable - 194Q Purchase of Goods", None),
                ("210306", "TDS Payable - 194A Interest", None),
                ("210307", "TDS Payable - 192 Salary", None),
            ]),
            ("210400", "Statutory Dues", [
                ("210401", "PF Payable", None),
                ("210402", "ESI Payable", None),
                ("210403", "Professional Tax Payable", None),
            ]),
            ("210500", "Provisions", [
                ("210501", "Salary Payable", None),
                ("210502", "Provision for Expenses", None),
                ("210503", "Audit Fees Payable", None),
                ("210504", "Provision for Income Tax", None),
            ]),
            ("210600", "Other Current Liabilities", [
                ("210601", "Advance from Customers", None),
                ("210602", "Goods Received Not Invoiced (Allocation)", None),
            ]),
        ]),
    ]),
    ("EQUITY", [
        ("300000", "Share Capital", [
            ("300100", "Equity Share Capital", None),
        ]),
        ("310000", "Reserves and Surplus", [
            ("310100", "General Reserve", None),
            ("310200", "Retained Earnings", None),
            ("310300", "Opening Balance Difference", None),
        ]),
    ]),
    ("REVENUE", [
        ("400000", "Revenue from Operations", [
            ("400100", "Domestic Sales - Intra State", None),
            ("400200", "Domestic Sales - Inter State", None),
            ("400300", "Export Sales", None),
            ("400400", "Sales Returns", None),
            ("400500", "Scrap Sales", None),
        ]),
        ("410000", "Other Income", [
            ("410100", "Interest Income", None),
            ("410200", "Discount Received", None),
            ("410300", "Foreign Exchange Gain", None),
        ]),
    ]),
    ("COST_OF_SALES", [
        ("500000", "Cost of Materials", [
            ("500100", "Purchase - Raw Material", None),
            ("500200", "Purchase - Import", None),
            ("500300", "Freight Inward", None),
            ("500400", "Cost of Goods Sold", None),
            ("500500", "Inventory Price Difference", None),
            ("500600", "Inventory Variance", None),
        ]),
    ]),
    ("EXPENSES", [
        ("510000", "Employee Costs", [
            ("510100", "Salaries and Wages", None),
            ("510200", "Contribution to PF", None),
            ("510300", "Contribution to ESI", None),
            ("510400", "Staff Welfare", None),
        ]),
        ("520000", "Manufacturing Expenses", [
            ("520100", "Power and Fuel", None),
            ("520200", "Repairs - Machinery", None),
            ("520300", "Job Work Charges", None),
            ("520400", "Consumables", None),
        ]),
        ("530000", "Administrative and Selling Expenses", [
            ("530100", "Rent - Building", None),
            ("530200", "Rent - Machinery Hire", None),
            ("530300", "Professional Fees", None),
            ("530400", "Technical Service Fees", None),
            ("530500", "Audit Fees", None),
            ("530600", "Commission on Sales", None),
            ("530700", "Telephone and Internet", None),
            ("530800", "Printing and Stationery", None),
            ("530900", "Travelling and Conveyance", None),
            ("531000", "Freight Outward", None),
            ("531100", "Advertisement", None),
            ("531200", "Insurance", None),
            ("531300", "Bank Charges", None),
        ]),
        ("540000", "Finance Costs", [
            ("540100", "Interest on Term Loan", None),
            ("540200", "Interest on Unsecured Loans", None),
        ]),
        ("550000", "Depreciation and Tax", [
            ("550100", "Depreciation", None),
            ("550200", "Income Tax Expense", None),
        ]),
    ]),
]
CONTROL_ACCOUNTS = {"110201", "110202", "210101", "210102"}  # BP control accounts - lock manual JEs


def flatten():
    rows = []
    def walk(nodes, parent, drawer, level):
        for code, name, kids in nodes:
            rows.append({"Code": code, "Name": name, "Parent": parent, "Drawer": drawer, "Level": level,
                         "Title": kids is not None, "AccountType": ACCT_TYPE.get(drawer, "at_Other")})
            if kids:
                walk(kids, code, drawer, level + 1)
    for drawer, nodes in COA:
        walk(nodes, None, drawer, 2)
    return rows


ACCOUNTS = flatten()
NAME = {a["Code"]: a["Name"] for a in ACCOUNTS}
assert len(NAME) == len(ACCOUNTS), "duplicate account code"

# ---------------------------------------------------------------- GST tax codes (dummy)
# GST 2.0 (w.e.f. 22-Sep-2025): slabs 0 / 5 / 18 / 40. 12% and 28% kept as LEGACY for old documents.
GST = []
for rate, note in [(0, "Nil rated / exempt"), (5, "Current slab"), (18, "Current slab"), (40, "Current slab - demerit goods"),
                   (12, "LEGACY - pre 22-Sep-2025"), (28, "LEGACY - pre 22-Sep-2025")]:
    half = rate / 2
    GST.append({"Code": f"CSG{rate}", "Description": f"CGST {half:g}% + SGST {half:g}% (intra-state)", "Rate": rate,
                "Type": "Intra-state", "Note": note,
                "Lines": [("CGST", half, "210201", "110401"), ("SGST", half, "210202", "110402")]})
    GST.append({"Code": f"IGST{rate}", "Description": f"IGST {rate}% (inter-state)", "Rate": rate,
                "Type": "Inter-state", "Note": note, "Lines": [("IGST", rate, "210203", "110403")]})

# ---------------------------------------------------------------- TDS / withholding tax codes (dummy)
# B1 withholding tax codes are at most 4 characters. Base = taxable value excluding GST (CBDT Circular 23/2017).
TDS = [
    # code, name, section, payee, rate %, threshold (INR), account, typical expense account
    ("C1", "TDS Contractor - Individual/HUF 1%", "194C", "Individual / HUF", 1, "30,000 single / 1,00,000 p.a.", "210301", "520300"),
    ("C2", "TDS Contractor - Others 2%", "194C", "Company / Firm / Others", 2, "30,000 single / 1,00,000 p.a.", "210301", "520300"),
    ("J10", "TDS Professional Fees 10%", "194J", "All", 10, "50,000 p.a.", "210302", "530300"),
    ("J2", "TDS Technical Services 2%", "194J", "All", 2, "50,000 p.a.", "210302", "530400"),
    ("I10", "TDS Rent - Land/Building 10%", "194I", "All", 10, "50,000 per month", "210303", "530100"),
    ("I2", "TDS Rent - Plant/Machinery 2%", "194I", "All", 2, "50,000 per month", "210303", "530200"),
    ("H2", "TDS Commission/Brokerage 2%", "194H", "All", 2, "20,000 p.a.", "210304", "530600"),
    ("Q01", "TDS Purchase of Goods 0.1%", "194Q", "All", 0.1, "Purchases above 50 lakh p.a.", "210305", "500100"),
    ("A10", "TDS Interest (non-bank) 10%", "194A", "All", 10, "10,000 p.a.", "210306", "540200"),
]
assert all(len(t[0]) <= 4 for t in TDS)

# ---------------------------------------------------------------- G/L account determination (suggested)
GL_DET = [
    ("Sales", "Domestic Accounts Receivable", "110201"), ("Sales", "Foreign Accounts Receivable", "110202"),
    ("Sales", "Revenue Account", "400100"), ("Sales", "Sales Returns", "400400"),
    ("Sales", "Exchange Rate Gains", "410300"),
    ("Purchasing", "Domestic Accounts Payable", "210101"), ("Purchasing", "Foreign Accounts Payable", "210102"),
    ("Purchasing", "Expense Account", "500100"), ("Purchasing", "Freight Inward", "500300"),
    ("Inventory", "Inventory Account", "110101"), ("Inventory", "Cost of Goods Sold", "500400"),
    ("Inventory", "Allocation Account", "210602"), ("Inventory", "Price Difference Account", "500500"),
    ("Inventory", "Variance Account", "500600"), ("Inventory", "WIP Inventory Account", "110102"),
    ("General", "Opening Balance Account", "310300"), ("General", "Bank Charges", "531300"),
    ("General", "Petty Cash", "110302"), ("General", "Period-end Closing Account", "310200"),
]
for _, _, acct in GL_DET:
    assert acct in NAME, acct
for g in GST:
    for line in g["Lines"]:
        assert line[2] in NAME and line[3] in NAME
for t in TDS:
    assert t[6] in NAME and t[7] in NAME

# ---------------------------------------------------------------- DTW files (one per level, load in order)
DTW_FIELDS = ["RecordKey", "Code", "Name", "FatherAccountKey", "ActiveAccount", "AccountType", "AcctCurrency",
              "LockManualTransaction"]
DTW_DB = ["RecordKey", "AcctCode", "AcctName", "FatherNum", "Postable", "ActType", "ActCurr", "LocManTran"]
for f in DTW.glob("COA_L*.csv"):
    f.unlink()
for level in sorted({a["Level"] for a in ACCOUNTS}):
    with open(DTW / f"COA_L{level}.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(DTW_FIELDS)
        w.writerow(DTW_DB)
        for i, a in enumerate((a for a in ACCOUNTS if a["Level"] == level), 1):
            w.writerow([i, a["Code"], a["Name"], a["Parent"] or DRAWERS[a["Drawer"]],
                        "tNO" if a["Title"] else "tYES", a["AccountType"],
                        "" if a["Title"] else CURRENCY,
                        "tYES" if a["Code"] in CONTROL_ACCOUNTS else "tNO"])

# ---------------------------------------------------------------- Excel workbook
F = "Arial"
HEAD = PatternFill("solid", fgColor="1F4E78")
TITLE_FILL = PatternFill("solid", fgColor="DDEBF7")
INPUT = PatternFill("solid", fgColor="FFFF00")
LEGACY = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()


def sheet(title, headers, widths, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BOX
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=F)
            c.border = BOX
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


readme = wb.active
readme.title = "Read Me"
lines = [
    ("DUMMY Finance Setup - SAP Business One (India)", True),
    ("All data in this workbook is DUMMY test data, not AKE's real Chart of Accounts.", False),
    ("", False),
    ("Load order (no Service Layer)", True),
    ("1. Create the test company with a User-Defined Chart of Accounts.", False),
    ("2. Fill the yellow cells on 'Drawer Mapping' (SQL shown there), copy the codes into DRAWERS in "
     "tools/build_dummy_setup.py and re-run it.", False),
    ("3. DTW > Import > Chart of Accounts: import data/dummy/dtw/COA_L2.csv, then COA_L3.csv, then COA_L4.csv "
     "(parents must exist first). Run a simulation first.", False),
    ("4. GST: Administration > Setup > Financials > Tax > Tax Codes - key in the codes on 'GST Tax Codes' "
     "(one line per component, accounts as shown). Then set Tax Code Determination: same state = CSG*, other state = IGST*.", False),
    ("5. TDS: Administration > Setup > Financials > Tax > Withholding Tax - key in the codes on 'TDS Codes'. "
     "Then tick 'Subject to Withholding Tax' on the vendor and assign the code.", False),
    ("6. G/L Account Determination: set the defaults listed on 'GL Determination'.", False),
    ("7. Test: one A/R + one A/P invoice for an intra-state and an inter-state BP, and one A/P invoice with TDS.", False),
    ("", False),
    ("Notes", True),
    ("- GST rates follow GST 2.0 (w.e.f. 22-Sep-2025: 0/5/18/40). 12% and 28% codes are LEGACY, kept for old "
     "documents and credit notes.", False),
    ("- TDS sections are shown with the familiar Income-tax Act 1961 numbers. From 1-Apr-2026 the Income-tax Act 2025 "
     "renumbers them - confirm section references, rates and thresholds with the CA before go-live.", False),
    ("- B1 withholding tax codes allow at most 4 characters (so 'C1', not '194C1').", False),
    ("- DTW columns use DI API enum values (tYES/tNO, at_Revenues/at_Expenses/at_Other). Compare the header row with "
     "the ChartOfAccounts template shipped with your DTW version before importing.", False),
    (f"- Accounts: {len(ACCOUNTS)} ({sum(a['Title'] for a in ACCOUNTS)} titles, "
     f"{sum(not a['Title'] for a in ACCOUNTS)} postable) | GST codes: {len(GST)} | TDS codes: {len(TDS)}", False),
]
for i, (text, bold) in enumerate(lines, 1):
    c = readme.cell(row=i, column=1, value=text)
    c.font = Font(name=F, bold=bold, size=14 if i == 1 else 10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
readme.column_dimensions["A"].width = 120

coa_ws = sheet("Chart of Accounts",
               ["Account Code", "Account Name", "Level", "Title / Postable", "Parent Code", "Parent Name",
                "Drawer", "Account Type", "Currency", "Control Account"],
               [14, 46, 7, 15, 13, 34, 16, 14, 10, 10],
               [[a["Code"], ("    " * (a["Level"] - 2)) + a["Name"], a["Level"], "Title" if a["Title"] else "Postable",
                 a["Parent"] or "(drawer)", NAME.get(a["Parent"], DRAWERS[a["Drawer"]] if not a["Parent"] else ""),
                 a["Drawer"], a["AccountType"], "" if a["Title"] else CURRENCY,
                 "Yes" if a["Code"] in CONTROL_ACCOUNTS else ""] for a in ACCOUNTS])
for row in coa_ws.iter_rows(min_row=2):
    if row[3].value == "Title":
        for c in row:
            c.fill = TITLE_FILL
            c.font = Font(name=F, bold=True)

gst_rows = []
for g in GST:
    for comp, r, out_acct, in_acct in g["Lines"]:
        gst_rows.append([g["Code"], g["Description"], g["Type"], g["Rate"] / 100, comp, r / 100,
                         out_acct, NAME[out_acct], in_acct, NAME[in_acct], g["Note"]])
gst_ws = sheet("GST Tax Codes",
               ["Tax Code", "Description", "Supply Type", "Total Rate", "Tax Type", "Component Rate",
                "Sales (Output) Account", "Output Account Name", "Purchase (Input) Account", "Input Account Name", "Status"],
               [10, 38, 12, 10, 9, 11, 14, 18, 14, 18, 26], gst_rows)
for row in gst_ws.iter_rows(min_row=2):
    row[3].number_format = row[5].number_format = "0.0%"
    if row[10].value.startswith("LEGACY"):
        for c in row:
            c.fill = LEGACY

tds_ws = sheet("TDS Codes",
               ["WT Code", "Name", "Section", "Payee Type", "Rate", "Base Type", "Base Amount %", "Category",
                "Threshold (INR)", "TDS Payable Account", "Account Name", "Typical Expense Account", "Expense Name"],
               [9, 34, 9, 22, 8, 22, 10, 10, 26, 12, 34, 12, 26],
               [[c, n, s, p, r / 100, "Net (excluding GST)", 1, "Invoice", th, a, NAME[a], e, NAME[e]]
                for c, n, s, p, r, th, a, e in TDS])
for row in tds_ws.iter_rows(min_row=2):
    row[4].number_format = "0.0#%"
    row[6].number_format = "0%"

sheet("GL Determination", ["Tab", "Determination Field", "Account Code", "Account Name"], [12, 32, 13, 40],
      [[t, f, a, NAME[a]] for t, f, a in GL_DET])

dm = sheet("Drawer Mapping", ["Drawer Key (script)", "Drawer Code in your company", "Used by top-level titles"],
           [20, 36, 60],
           [[k, "", ", ".join(f"{a['Code']} {a['Name']}" for a in ACCOUNTS if a["Drawer"] == k and a["Level"] == 2)]
            for k in DRAWERS])
for row in dm.iter_rows(min_row=2):
    row[1].fill = INPUT
n = dm.max_row + 2
dm.cell(row=n, column=1, value="Fill the yellow cells with: "
        'SELECT "AcctCode", "AcctName" FROM OACT WHERE "Levels" = 1 ORDER BY "AcctCode";').font = Font(name=F, italic=True)
dm.cell(row=n + 1, column=1, value="Then copy them into DRAWERS in tools/build_dummy_setup.py and re-run the script "
        "to regenerate the DTW files.").font = Font(name=F, italic=True)

wb.save(OUT / "Dummy_Finance_Setup.xlsx")
print(f"Accounts {len(ACCOUNTS)}  GST codes {len(GST)}  TDS codes {len(TDS)}")
print(f"Wrote {OUT / 'Dummy_Finance_Setup.xlsx'} and {sorted(p.name for p in DTW.glob('*.csv'))}")
