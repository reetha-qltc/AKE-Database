"""Build import files for AKE's REAL Chart of Accounts, GST tax codes and TDS codes - no Service Layer needed.

Run:  python tools/build_ake_import.py
Input:  data/01_chart_of_accounts.json (built by tools/build_data.py from AKE's COA export)
Output (data/ake_import/):
  AKE_Finance_Setup.xlsx       - COA, GST tax codes, TDS codes, drawer mapping, findings
  dtw/AKE_COA_L2.csv ... L5    - Chart of Accounts in DTW format; import in level order on the SAP B1 machine

Set DRAWERS to the target company's Level-1 drawer codes, then re-run:
  SELECT "AcctCode", "AcctName" FROM OACT WHERE "Levels" = 1 ORDER BY "AcctCode";
"""
import csv, json, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "ake_import"
DTW = OUT / "dtw"
DTW.mkdir(parents=True, exist_ok=True)

DRAWERS = {  # AKE section -> drawer code in the target company. EDIT before importing.
    "Asset": "<Assets drawer code>",
    "Liability": "<Liabilities drawer code>",
    "Equity": "<Capital and Reserves drawer code>",
    "Revenue": "<Revenue drawer code>",
    "Expenditure": "<Expenses drawer code>",
}
CURRENCY = "INR"

ACCOUNTS = json.loads((ROOT / "data" / "01_chart_of_accounts.json").read_text(encoding="utf-8"))
NAME = {a["Code"]: a["Name"] for a in ACCOUNTS}
CONTROL = {b["ControlAccount"] for b in json.loads((ROOT / "data" / "06_business_partners.json").read_text(encoding="utf-8"))}
for a in ACCOUNTS:
    a["Level"] = a["Depth"] + 2  # drawer is level 1

# ---------------------------------------------------------------- GST (AKE's own ledgers, as in the setup checklist)
# B1 tax codes are max 8 characters, so 'CGST+SGST18' becomes 'CSG18'.
OUT_C, OUT_S, OUT_I = "3002-03-24", "3002-03-25", "3002-03-22"
IN_C, IN_S, IN_I = "5002-05-02-02", "5002-05-02-01", "5002-05-02-03"
GST = []
for rate, note in [(0, "Nil rated / exempt"), (5, "Current slab"), (18, "Current slab"), (40, "Current slab - demerit goods"),
                   (12, "LEGACY - pre 22-Sep-2025 (AKE items still at 12%)"), (28, "LEGACY - pre 22-Sep-2025")]:
    h = rate / 2
    GST.append((f"CSG{rate}", f"CGST {h:g}% + SGST {h:g}%", "Intra-state", rate, note,
                [("CGST", h, OUT_C, IN_C), ("SGST", h, OUT_S, IN_S)]))
    GST.append((f"IGST{rate}", f"IGST {rate}%", "Inter-state", rate, note, [("IGST", rate, OUT_I, IN_I)]))

# ---------------------------------------------------------------- TDS (AKE's own ledgers). B1 WT codes are max 4 characters.
TDS = [  # code, old code in checklist, name, section, payee, rate, threshold, account
    ("Q01", "194Q", "TDS on Purchase of Goods 0.1%", "194Q", "All", 0.1, "Purchases above 50 lakh p.a.", "3002-03-23"),
    ("C1", "194C1", "TDS Contractor - Individual/HUF 1%", "194C", "Individual / HUF", 1, "30,000 single / 1,00,000 p.a.", "3002-03-18"),
    ("C2", "194C2", "TDS Contractor - Others 2%", "194C", "Company / Firm / Others", 2, "30,000 single / 1,00,000 p.a.", "3002-03-14"),
    ("J10", "194J", "TDS Professional Fees 10%", "194J", "All", 10, "50,000 p.a.", "3002-03-16"),
    ("I10", "194I", "TDS on Rent 10%", "194I", "All", 10, "50,000 per month", "3002-03-15"),
    ("A10", "194A", "TDS on Interest 10%", "194A", "All", 10, "10,000 p.a.", "3002-03-26"),
]
VENDORS = {"Q01": "V0001, V0002", "C2": "V0007, V0008", "J10": "V0009"}
for g in GST:
    for line in g[5]:
        assert line[2] in NAME and line[3] in NAME, line
for t in TDS:
    assert len(t[0]) <= 4 and t[7] in NAME, t

# Points the client must decide - found while mapping (not changed in the data)
FINDINGS = [
    ("GST output ledgers exist twice", "3002-03-22/24/25 (IGST/CGST/SGST Payable) and 3002-09-01/02 + 3003-09-03 "
     "(... Output Payable). Tax codes use 3002-03-xx as per the checklist - AKE to confirm which set to keep."),
    ("Tax code names too long", "B1 tax codes allow 8 characters: 'CGST+SGST18' renamed to 'CSG18'."),
    ("TDS code names too long", "B1 withholding tax codes allow 4 characters: '194C1' renamed to 'C1' etc. (column 'Old Code')."),
    ("12% GST slab removed", "GST 2.0 (22-Sep-2025) moved most 12% goods to 5%. AKE items mapped to 12% need HSN review."),
    ("RCM ledgers present", "5002-05-03-01..03 and 3003-09-06..08 exist for reverse charge - RCM tax codes not created yet."),
    ("TDS section numbers", "Income-tax Act 2025 renumbers TDS sections from 1-Apr-2026 - CA to confirm sections, rates, thresholds."),
]

# ---------------------------------------------------------------- DTW files
FIELDS = ["RecordKey", "Code", "Name", "FatherAccountKey", "ActiveAccount", "AccountType", "AcctCurrency", "LockManualTransaction"]
DBCOLS = ["RecordKey", "AcctCode", "AcctName", "FatherNum", "Postable", "ActType", "ActCurr", "LocManTran"]
for f in DTW.glob("AKE_COA_L*.csv"):
    f.unlink()
levels = sorted({a["Level"] for a in ACCOUNTS})
for lv in levels:
    with open(DTW / f"AKE_COA_L{lv}.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(FIELDS)
        w.writerow(DBCOLS)
        for i, a in enumerate((a for a in ACCOUNTS if a["Level"] == lv), 1):
            w.writerow([i, a["Code"], a["Name"], a["Parent"] or DRAWERS[a["Section"]], "tNO" if a["Title"] else "tYES",
                        a["AccountType"], "" if a["Title"] else CURRENCY, "tYES" if a["Code"] in CONTROL else "tNO"])

# ---------------------------------------------------------------- Excel
F = "Arial"
HEAD, TFILL = PatternFill("solid", fgColor="1F4E78"), PatternFill("solid", fgColor="DDEBF7")
YEL, GREY, ORANGE = PatternFill("solid", fgColor="FFFF00"), PatternFill("solid", fgColor="F2F2F2"), PatternFill("solid", fgColor="FCE4D6")
s = Side(style="thin", color="BFBFBF")
BOX = Border(left=s, right=s, top=s, bottom=s)
wb = Workbook()


def sheet(title, headers, widths, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.border = Font(name=F, bold=True, color="FFFFFF"), HEAD, BOX
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.border, c.alignment = Font(name=F), BOX, Alignment(vertical="top", wrap_text=True)
    for i, wd in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


rd = wb.active
rd.title = "Read Me"
n_t = sum(a["Title"] for a in ACCOUNTS)
text = [
    ("AKE Finance Setup - Chart of Accounts, GST, TDS (SAP Business One, no Service Layer)", True),
    (f"Chart of Accounts: {len(ACCOUNTS)} accounts ({n_t} titles, {len(ACCOUNTS) - n_t} postable), levels {levels[0]}-{levels[-1]}. "
     f"GST tax codes: {len(GST)}. TDS codes: {len(TDS)}.", False),
    ("", False),
    ("Steps on the SAP B1 machine (SSO-WINDOWS1)", True),
    ("1. Company must be created with a User-Defined Chart of Accounts.", False),
    ("2. Run the SQL on 'Drawer Mapping', fill the yellow cells, put the codes in DRAWERS in tools/build_ake_import.py, re-run.", False),
    ("3. Copy data/ake_import/dtw/*.csv to the SAP machine. DTW > Import > Business Object 'Chart of Accounts'. "
     + " -> ".join(f"AKE_COA_L{lv}.csv" for lv in levels) + " (parents first). Simulate first, then import.", False),
    ("4. Tax Codes: Administration > Setup > Financials > Tax > Tax Codes - key in 'GST Tax Codes' (one row = one line).", False),
    ("   Tax Code Determination: Karnataka -> Karnataka = CSG*, other state = IGST*.", False),
    ("5. Withholding Tax: Administration > Setup > Financials > Tax > Withholding Tax - key in 'TDS Codes', then "
     "tick Subject to Withholding Tax on the vendors listed.", False),
    ("6. Test: A/R + A/P invoice intra-state and inter-state; A/P invoice with TDS. Check the journal entry accounts.", False),
    ("", False),
    ("Read the 'Findings' sheet - those points need AKE's decision.", True),
]
for i, (t, b) in enumerate(text, 1):
    c = rd.cell(row=i, column=1, value=t)
    c.font, c.alignment = Font(name=F, bold=b, size=14 if i == 1 else 10), Alignment(wrap_text=True)
rd.column_dimensions["A"].width = 130

cw = sheet("Chart of Accounts", ["Account Code", "Account Name", "Level", "Title / Postable", "Parent Code", "Parent Name",
                                 "Section", "Account Type", "Control Account", "Name blank in AKE source"],
           [15, 50, 7, 14, 15, 36, 12, 13, 10, 12],
           [[a["Code"], "    " * a["Depth"] + a["Name"], a["Level"], "Title" if a["Title"] else "Postable",
             a["Parent"] or "(drawer)", NAME.get(a["Parent"], ""), a["Section"], a["AccountType"],
             "Yes" if a["Code"] in CONTROL else "", "Yes" if a["NameBlankInSource"] else ""] for a in ACCOUNTS])
for row in cw.iter_rows(min_row=2):
    if row[3].value == "Title":
        for c in row:
            c.fill, c.font = TFILL, Font(name=F, bold=True)
    if row[9].value:
        row[1].fill = ORANGE

gw = sheet("GST Tax Codes", ["Tax Code", "Description", "Supply Type", "Total Rate", "Tax Type", "Line Rate",
                             "Output Account", "Output Account Name", "Input Account", "Input Account Name", "Status"],
           [10, 24, 12, 10, 9, 9, 14, 22, 15, 22, 34],
           [[c, d, ty, r / 100, comp, lr / 100, oa, NAME[oa], ia, NAME[ia], note]
            for c, d, ty, r, note, lines in GST for comp, lr, oa, ia in lines])
for row in gw.iter_rows(min_row=2):
    row[3].number_format = row[5].number_format = "0.0%"
    if row[10].value.startswith("LEGACY"):
        for c in row:
            c.fill = GREY

tw = sheet("TDS Codes", ["WT Code", "Old Code", "Name", "Section", "Payee Type", "Rate", "Base Type", "Base Amount %",
                         "Category", "Threshold (INR)", "TDS Payable Account", "Account Name", "Assign to Vendors"],
           [9, 9, 34, 9, 22, 8, 20, 10, 10, 26, 14, 36, 16],
           [[c, old, nm, sec, p, r / 100, "Net (excluding GST)", 1, "Invoice", th, acc, NAME[acc], VENDORS.get(c, "")]
            for c, old, nm, sec, p, r, th, acc in TDS])
for row in tw.iter_rows(min_row=2):
    row[5].number_format, row[7].number_format = "0.0#%", "0%"

dm = sheet("Drawer Mapping", ["AKE Section", "Drawer Code in target company", "Top-level AKE titles under it"], [14, 34, 90],
           [[k, "", ", ".join(f"{a['Code']} {a['Name']}" for a in ACCOUNTS if a["Section"] == k and a["Depth"] == 0)]
            for k in DRAWERS])
for row in dm.iter_rows(min_row=2):
    row[1].fill = YEL
dm.cell(row=dm.max_row + 2, column=1,
        value='Fill yellow cells from: SELECT "AcctCode", "AcctName" FROM OACT WHERE "Levels" = 1 ORDER BY "AcctCode";').font = Font(name=F, italic=True)

sheet("Findings", ["Topic", "Detail"], [28, 110], FINDINGS)
wb.save(OUT / "AKE_Finance_Setup.xlsx")
print(f"COA {len(ACCOUNTS)} ({n_t} titles)  levels {levels}  GST {len(GST)}  TDS {len(TDS)}")
print("Wrote", OUT / "AKE_Finance_Setup.xlsx", sorted(p.name for p in DTW.glob("*.csv")))
